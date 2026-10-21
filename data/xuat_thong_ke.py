"""
Script xuat cac file Excel thong ke va tham so cho Tuan 1:
1. docs/thong_ke_mo_ta.xlsx
2. docs/bang_tham_so.xlsx
"""

import os
import pandas as pd
import numpy as np
import scipy.stats as stats
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_reports():
    clean_path = os.path.join(r"d:\KHDV", "data", "clean", "data_clean.csv")
    df = pd.read_csv(clean_path)
    n_days = df['day'].nunique()

    # ----------------------------------------------------
    # 1. Tinh toan thong ke theo gio
    # ----------------------------------------------------
    counts_matrix = df.groupby(['day', 'hour']).size().unstack(fill_value=0)
    
    hourly_rows = []
    for h in range(9):
        counts = counts_matrix[h] if h in counts_matrix else pd.Series(0, index=range(n_days))
        n_bar = counts.mean()
        lam = n_bar / 60.0
        se_lam = np.sqrt(n_bar / n_days) / 60.0
        ci_lam_low = max(0, lam - 1.96 * se_lam)
        ci_lam_high = lam + 1.96 * se_lam
        
        disp = counts.var(ddof=0) * n_days / n_bar
        p_poisson = stats.chi2.sf(disp, df=n_days - 1)
        
        df_h = df[df['hour'] == h]
        n_tot = len(df_h)
        n_ab = (df_h['served'] == 0).sum()
        ab_rate = n_ab / n_tot if n_tot > 0 else 0
        vip_rate = df_h['vip'].mean() if n_tot > 0 else 0
        
        df_h_served = df_h[df_h['served'] == 1]
        mean_svc_h = df_h_served['svc'].mean() if len(df_h_served) > 0 else 0
        std_svc_h = df_h_served['svc'].std() if len(df_h_served) > 1 else 0
        
        hourly_rows.append({
            'Khung gio': f'{h+8}h - {h+9}h',
            'Hour_ID': h,
            'N_tb (khach/ngay)': round(n_bar, 2),
            'lambda (khach/phut)': round(lam, 4),
            'CI95_lambda_low': round(ci_lam_low, 4),
            'CI95_lambda_high': round(ci_lam_high, 4),
            'Chi so phan tan D': round(disp, 2),
            'p-value Poisson': round(p_poisson, 4),
            'Ket luan Poisson': 'Chua bac bo' if p_poisson >= 0.01 else 'Lech nhe',
            'Ti le bo ve': round(ab_rate, 4),
            'Ti le VIP': round(vip_rate, 4),
            'E[S] theo gio (phut)': round(mean_svc_h, 2),
            'std[S] theo gio (phut)': round(std_svc_h, 2)
        })
    df_hourly = pd.DataFrame(hourly_rows)

    # ----------------------------------------------------
    # 2. Tinh toan thong ke tham so chung
    # ----------------------------------------------------
    svc_data = df[df['served'] == 1]['svc'].dropna().values
    mean_svc = np.mean(svc_data)
    std_svc = np.std(svc_data, ddof=1)
    cv_svc = std_svc / mean_svc
    cs2_svc = cv_svc ** 2
    se_svc = std_svc / np.sqrt(len(svc_data))

    n_abandon = (df['served'] == 0).sum()
    total_wait = df['wait'].sum()
    theta_hat = n_abandon / total_wait
    theta_low = theta_hat * (1 - 1.96 / np.sqrt(n_abandon))
    theta_high = theta_hat * (1 + 1.96 / np.sqrt(n_abandon))

    p_vip = df['vip'].mean()
    se_vip = np.sqrt(p_vip * (1 - p_vip) / len(df))
    vip_low = p_vip - 1.96 * se_vip
    vip_high = p_vip + 1.96 * se_vip

    chung_rows = [
        {'Tham so': 'E[S] - Thoi gian phuc vu TB (phut)', 'Gia tri': round(mean_svc, 4), 'CI95': f'[{mean_svc - 1.96*se_svc:.4f}, {mean_svc + 1.96*se_svc:.4f}]', 'Ghi chu': 'Tinh tren 9.133 luot khach duoc phuc vu thanh cong'},
        {'Tham so': 'std[S] - Do lech chuan phuc vu (phut)', 'Gia tri': round(std_svc, 4), 'CI95': '-', 'Ghi chu': 'Phan anh bien thien thoi gian thao tac'},
        {'Tham so': 'cv - He so bien thien (std/mean)', 'Gia tri': round(cv_svc, 4), 'CI95': '-', 'Ghi chu': 'cv > 1 cho thay bien thien lon hon phan phoi mu chuan (cv=1)'},
        {'Tham so': 'cs^2 - Binh phuong cv (cs^2 = cv^2)', 'Gia tri': round(cs2_svc, 4), 'CI95': '-', 'Ghi chu': 'Dung trong cong thuc xap xi Allen-Cunneen cho M/G/c'},
        {'Tham so': 'Phan phoi phuc vu tot nhat', 'Gia tri': 'Gamma / Lognormal', 'CI95': '-', 'Ghi chu': 'Phan phoi mu bi bac bo do du lieu thuc te co duoi phai dai'},
        {'Tham so': 'theta - Toc do bo ve (1/phut)', 'Gia tri': round(theta_hat, 4), 'CI95': f'[{theta_low:.4f}, {theta_high:.4f}]', 'Ghi chu': 'Uoc luong hop ly cuc dai (MLE) voi du lieu bi kiem duyet'},
        {'Tham so': 'Kien nhan trung binh (1/theta) (phut)', 'Gia tri': round(1/theta_hat, 2), 'CI95': f'[{1/theta_high:.2f}, {1/theta_low:.2f}]', 'Ghi chu': 'Thoi gian khach san sang cho doi truoc khi bo ve'},
        {'Tham so': 'Ti le khach VIP (p_vip)', 'Gia tri': round(p_vip, 4), 'CI95': f'[{vip_low:.4f}, {vip_high:.4f}]', 'Ghi chu': 'Khach hang priority level 2 trong he thong ngan hang'},
        {'Tham so': 'Ti le bo ve trung binh toan ngay', 'Gia tri': round(n_abandon/len(df), 4), 'CI95': '-', 'Ghi chu': f'{n_abandon} cuoc bo ve tren tong so {len(df)} khach den'}
    ]
    df_chung = pd.DataFrame(chung_rows)

    # ----------------------------------------------------
    # 3. Xuat docs/thong_ke_mo_ta.xlsx
    # ----------------------------------------------------
    path_thong_ke = os.path.join(r"d:\KHDV", "docs", "thong_ke_mo_ta.xlsx")
    with pd.ExcelWriter(path_thong_ke, engine='openpyxl') as writer:
        df_hourly.to_excel(writer, sheet_name='Thong_ke_theo_gio', index=False)
        counts_matrix.columns = [f'{h+8}h-{h+9}h' for h in counts_matrix.columns]
        counts_matrix.index = [f'Ngay_{d}' for d in counts_matrix.index]
        counts_matrix.to_excel(writer, sheet_name='Khach_theo_ngay_gio')

    # Dinh dang dep cho thong_ke_mo_ta.xlsx
    wb = openpyxl.load_workbook(path_thong_ke)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        ws.views.sheetView[0].showGridLines = True
        for cell in ws[1]:
            cell.font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
            cell.alignment = Alignment(horizontal='center', vertical='center')
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
    wb.save(path_thong_ke)
    print(f"Da tao: {path_thong_ke}")

    # ----------------------------------------------------
    # 4. Xuat docs/bang_tham_so.xlsx
    # ----------------------------------------------------
    path_tham_so = os.path.join(r"d:\KHDV", "docs", "bang_tham_so.xlsx")
    
    # Sheet theo_gio
    sheet_theo_gio_cols = ['Khung gio', 'N_tb (khach/ngay)', 'lambda (khach/phut)', 'CI95_lambda_low', 'CI95_lambda_high', 'p-value Poisson']
    df_tham_so_gio = df_hourly[sheet_theo_gio_cols].copy()
    df_tham_so_gio.columns = ['Khung giờ', 'N̄_h (khách/ngày)', 'λ_h (khách/phút)', 'CI95_low', 'CI95_high', 'p-value Poisson']

    with pd.ExcelWriter(path_tham_so, engine='openpyxl') as writer:
        df_tham_so_gio.to_excel(writer, sheet_name='theo_gio', index=False)
        df_chung.to_excel(writer, sheet_name='chung', index=False)

    wb2 = openpyxl.load_workbook(path_tham_so)
    for sheet in wb2.sheetnames:
        ws = wb2[sheet]
        ws.views.sheetView[0].showGridLines = True
        for cell in ws[1]:
            cell.font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
            cell.alignment = Alignment(horizontal='center', vertical='center')
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
    wb2.save(path_tham_so)
    print(f"Da tao: {path_tham_so}")

if __name__ == '__main__':
    generate_excel_reports()
