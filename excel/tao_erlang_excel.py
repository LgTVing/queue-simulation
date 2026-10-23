"""
Script tao file Excel Erlang C: excel/erlang.xlsx
Sua loi triet de #NAME? bang cach dung ham tuong thich toan cau POISSON() thay vi POISSON.DIST().
Thong nhat o tham chieu dung 100% theo Muc 5.2 cua TUAN_1_HUONG_DAN.md.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

def build_erlang_excel():
    excel_path = os.path.join(BASE_DIR, "excel", "erlang.xlsx")
    os.makedirs(os.path.dirname(excel_path), exist_ok=True)
    
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Dinh dang
    font_title = Font(name='Segoe UI', size=13, bold=True, color='1F4E78')
    font_header = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    font_bold = Font(name='Segoe UI', size=11, bold=True)
    font_normal = Font(name='Segoe UI', size=11)

    fill_header_dark = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    fill_header_blue = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
    fill_input = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    fill_pass = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')

    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # =========================================================================
    # 1. SHEET Input (Chuan theo 5.2 TUAN_1_HUONG_DAN.md: B1 den B6)
    # =========================================================================
    ws_input = wb.create_sheet('Input')
    ws_input.views.sheetView[0].showGridLines = True

    input_rows = [
        ("A1", "λ (khách/phút)", "B1", 1.5, "C1", "khách/phút", "D1", "Tốc độ khách đến trung bình"),
        ("A2", "E[S] (phút)", "B2", 3.0, "C2", "phút", "D2", "Thời gian phục vụ trung bình"),
        ("A3", "μ = 1/E[S]", "B3", "=1/B2", "C3", "khách/phút", "D3", "Tốc độ phục vụ của 1 quầy"),
        ("A4", "a = λ*E[S] (Erlang)", "B4", "=B1*B2", "C4", "Erlang", "D4", "Cường độ tải giao thông"),
        ("A5", "Ngưỡng chờ t (phút)", "B5", 5.0, "C5", "phút", "D5", "Ngưỡng cam kết thời gian chờ SLA"),
        ("A6", "Mục tiêu P(chờ <= t)", "B6", 0.90, "C6", "90%", "D6", "Tỉ lệ tối thiểu khách chờ <= t")
    ]

    for r in input_rows:
        ws_input[r[0]] = r[1]
        ws_input[r[2]] = r[3]
        ws_input[r[4]] = r[5]
        ws_input[r[6]] = r[7]

    for row in range(1, 7):
        ws_input[f'A{row}'].font = font_bold
        ws_input[f'B{row}'].font = font_bold
        ws_input[f'B{row}'].fill = fill_input
        ws_input[f'B{row}'].alignment = Alignment(horizontal='center')
        ws_input[f'C{row}'].alignment = Alignment(horizontal='center')
        for col in ['A', 'B', 'C', 'D']:
            ws_input[f'{col}{row}'].border = border_thin

    ws_input['B1'].number_format = '0.00'
    ws_input['B2'].number_format = '0.00'
    ws_input['B3'].number_format = '0.0000'
    ws_input['B4'].number_format = '0.00'
    ws_input['B5'].number_format = '0.0'
    ws_input['B6'].number_format = '0.0%'

    # =========================================================================
    # 2. SHEET ErlangC
    # =========================================================================
    ws_erlang = wb.create_sheet('ErlangC')
    ws_erlang.views.sheetView[0].showGridLines = True

    ws_erlang['A1'] = "TÍNH TOÁN CÔNG THỨC ERLANG C VÀ SỐ QUẦY TỐI THIỂU"
    ws_erlang['A1'].font = font_title

    ws_erlang['A3'] = "Số quầy nhỏ nhất đạt SLA:"
    ws_erlang['A3'].font = font_bold
    ws_erlang['B3'] = '=INDEX(A10:A24, MATCH("ĐẠT", G10:G24, 0))'
    ws_erlang['B3'].font = Font(name='Segoe UI', size=14, bold=True, color='006100')
    ws_erlang['B3'].fill = fill_pass
    ws_erlang['B3'].alignment = Alignment(horizontal='center', vertical='center')

    ws_erlang['A4'] = "Hiệu suất sử dụng (ρ):"
    ws_erlang['A4'].font = font_bold
    ws_erlang['B4'] = '=VLOOKUP(B3, A10:G24, 2, FALSE)'
    ws_erlang['B4'].font = font_bold
    ws_erlang['B4'].number_format = '0.0%'
    ws_erlang['B4'].alignment = Alignment(horizontal='center', vertical='center')

    ws_erlang['A5'] = "Thời gian chờ TB Wq:"
    ws_erlang['A5'].font = font_bold
    ws_erlang['B5'] = '=VLOOKUP(B3, A10:G24, 4, FALSE)'
    ws_erlang['B5'].font = font_bold
    ws_erlang['B5'].number_format = '0.00 "phút"'
    ws_erlang['B5'].alignment = Alignment(horizontal='center', vertical='center')

    ws_erlang['A6'] = "P(Wq <= 5 phút):"
    ws_erlang['A6'].font = font_bold
    ws_erlang['B6'] = '=VLOOKUP(B3, A10:G24, 6, FALSE)'
    ws_erlang['B6'].font = Font(name='Segoe UI', size=11, bold=True, color='006100')
    ws_erlang['B6'].number_format = '0.0%'
    ws_erlang['B6'].alignment = Alignment(horizontal='center', vertical='center')

    headers_erlang = [
        "Số quầy c", "ρ = a/c", "Xác suất chờ C(c, a)", 
        "Wq (phút)", "Lq (khách)", "P(Wq <= t)", "Đạt mục tiêu?"
    ]
    for col_idx, h_text in enumerate(headers_erlang, start=1):
        cell = ws_erlang.cell(row=9, column=col_idx, value=h_text)
        cell.font = font_header
        cell.fill = fill_header_dark
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Dong 10 tro xuong (tuyet doi dung ham POISSON() thay cho POISSON.DIST() de khong bi loi #NAME?)
    for i in range(15):
        r = 10 + i
        if i == 0:
            ws_erlang[f'A{r}'] = '=INT(Input!$B$4)+1'
        else:
            ws_erlang[f'A{r}'] = f'=A{r-1}+1'
        
        ws_erlang[f'B{r}'] = f'=Input!$B$4/A{r}'
        # Cong thuc POISSON chuan tuong thich moi phien ban Excel
        ws_erlang[f'C{r}'] = f'=(POISSON(A{r},Input!$B$4,FALSE)/(1-B{r}))/((POISSON(A{r},Input!$B$4,FALSE)/(1-B{r}))+POISSON(A{r}-1,Input!$B$4,TRUE))'
        # Wq = C / (c*mu - lambda)
        ws_erlang[f'D{r}'] = f'=C{r}/(A{r}*Input!$B$3-Input!$B$1)'
        # Lq = lambda * Wq
        ws_erlang[f'E{r}'] = f'=Input!$B$1*D{r}'
        # P(Wq <= t) = 1 - C * exp(-(c*mu - lambda)*t)
        ws_erlang[f'F{r}'] = f'=1-C{r}*EXP(-(A{r}*Input!$B$3-Input!$B$1)*Input!$B$5)'
        # Dat muc tieu?
        ws_erlang[f'G{r}'] = f'=IF(F{r}>=Input!$B$6,"ĐẠT","KHÔNG")'

        ws_erlang[f'A{r}'].alignment = Alignment(horizontal='center')
        ws_erlang[f'A{r}'].font = font_bold
        ws_erlang[f'B{r}'].number_format = '0.0%'
        ws_erlang[f'B{r}'].alignment = Alignment(horizontal='center')
        ws_erlang[f'C{r}'].number_format = '0.0000'
        ws_erlang[f'C{r}'].alignment = Alignment(horizontal='center')
        ws_erlang[f'D{r}'].number_format = '0.00'
        ws_erlang[f'D{r}'].alignment = Alignment(horizontal='center')
        ws_erlang[f'E{r}'].number_format = '0.00'
        ws_erlang[f'E{r}'].alignment = Alignment(horizontal='center')
        ws_erlang[f'F{r}'].number_format = '0.0%'
        ws_erlang[f'F{r}'].alignment = Alignment(horizontal='center')
        ws_erlang[f'G{r}'].alignment = Alignment(horizontal='center')
        ws_erlang[f'G{r}'].font = font_bold

        for col_l in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
            ws_erlang[f'{col_l}{r}'].border = border_thin

    # =========================================================================
    # 3. SHEET Tra_c
    # =========================================================================
    ws_trac = wb.create_sheet('Tra_c')
    ws_trac.views.sheetView[0].showGridLines = True

    ws_trac['A1'] = "BẢNG TRA SỐ QUẦY TỐI THIỂU ĐẠT SLA 90% (t <= 5 phút)"
    ws_trac['A1'].font = font_title

    es_values = [2.0, 3.0, 4.0, 5.0]
    ws_trac['A3'] = "λ (khách/phút)"
    ws_trac['A3'].font = font_header
    ws_trac['A3'].fill = fill_header_dark
    ws_trac['A3'].alignment = Alignment(horizontal='center', vertical='center')

    for idx, es in enumerate(es_values, start=2):
        col_l = get_column_letter(idx)
        cell = ws_trac[f'{col_l}3']
        cell.value = f"E[S] = {es:.1f} phút (μ = {1/es:.2f})"
        cell.font = font_header
        cell.fill = fill_header_blue
        cell.alignment = Alignment(horizontal='center', vertical='center')

    import math
    def erlang_c_prob(c, a):
        if c <= a:
            return 1.0
        rho = a / c
        sum_k = sum((a**k)/math.factorial(k) for k in range(c))
        term_c = (a**c)/(math.factorial(c) * (1 - rho))
        return term_c / (sum_k + term_c)

    def find_min_c(lam, es, t=5.0, target=0.90):
        mu = 1.0 / es
        a = lam * es
        for c in range(int(a) + 1, int(a) + 50):
            C = erlang_c_prob(c, a)
            sl = 1.0 - C * math.exp(-(c * mu - lam) * t)
            if sl >= target:
                return c
        return int(a) + 1

    lambdas = [0.5, 0.75, 1.0, 1.25, 1.5, 1.75, 2.0, 2.25, 2.5, 2.75, 3.0]
    for r_idx, lam in enumerate(lambdas, start=4):
        ws_trac[f'A{r_idx}'] = lam
        ws_trac[f'A{r_idx}'].number_format = '0.00'
        ws_trac[f'A{r_idx}'].font = font_bold
        ws_trac[f'A{r_idx}'].alignment = Alignment(horizontal='center')
        ws_trac[f'A{r_idx}'].border = border_thin

        for c_idx, es in enumerate(es_values, start=2):
            col_l = get_column_letter(c_idx)
            c_min = find_min_c(lam, es)
            cell = ws_trac[f'{col_l}{r_idx}']
            cell.value = c_min
            cell.font = font_normal
            cell.alignment = Alignment(horizontal='center')
            cell.border = border_thin
            if es == 3.0 and lam == 1.5:
                cell.fill = fill_pass
                cell.font = font_bold

    # =========================================================================
    # 4. SHEET Allen_Cunneen
    # =========================================================================
    ws_ac = wb.create_sheet('Allen_Cunneen')
    ws_ac.views.sheetView[0].showGridLines = True

    ws_ac['A1'] = "XẤP XỈ ALLEN - CUNNEEN CHO HÀNG ĐỢI M/G/c"
    ws_ac['A1'].font = font_title
    ws_ac['A2'] = "Công thức: Wq(M/G/c) ≈ (1 + cs²) / 2 * Wq(M/M/c)"
    ws_ac['A2'].font = font_bold

    headers_ac = ["Số quầy c", "Wq M/M/c (cs²=1.0)", "Wq (cs²=0.5)", "Wq (cs²=1.2 - Thực tế)", "Wq (cs²=1.5)", "Tăng thời gian chờ do cs²>1"]
    for idx, h_text in enumerate(headers_ac, start=1):
        cell = ws_ac.cell(row=4, column=idx, value=h_text)
        cell.font = font_header
        cell.fill = fill_header_dark
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for i in range(10):
        r = 5 + i
        c_val = 5 + i
        ws_ac[f'A{r}'] = c_val
        ws_ac[f'A{r}'].alignment = Alignment(horizontal='center')
        ws_ac[f'A{r}'].font = font_bold
        
        ws_ac[f'B{r}'] = f'=ErlangC!D{10+i}'
        ws_ac[f'C{r}'] = f'=(1 + 0.5)/2 * B{r}'
        ws_ac[f'D{r}'] = f'=(1 + 1.20)/2 * B{r}'
        ws_ac[f'E{r}'] = f'=(1 + 1.5)/2 * B{r}'
        ws_ac[f'F{r}'] = f'=(D{r} - B{r})/B{r}'

        ws_ac[f'B{r}'].number_format = '0.00'
        ws_ac[f'C{r}'].number_format = '0.00'
        ws_ac[f'D{r}'].number_format = '0.00'
        ws_ac[f'E{r}'].number_format = '0.00'
        ws_ac[f'F{r}'].number_format = '+0.0%'

        for col_l in ['A', 'B', 'C', 'D', 'E', 'F']:
            ws_ac[f'{col_l}{r}'].border = border_thin
            ws_ac[f'{col_l}{r}'].alignment = Alignment(horizontal='center')

    # =========================================================================
    # 5. SHEET Quyet_dinh (tham so TRAIN tu results/final_v3/params.json)
    # =========================================================================
    import sys, json
    sys.path.insert(0, BASE_DIR)
    import b_theory as bt

    with open(os.path.join(BASE_DIR, "results", "final_v3", "params.json"), encoding="utf-8") as f:
        prm = json.load(f)
    train_lambdas = prm["rates_per_min"]
    es_train = prm["service"]["mean"]
    cs2_train = prm["service"]["cs2"]
    theta_train = prm["theta"]
    mu_train = 1.0 / es_train

    ws_qd = wb.create_sheet('Quyet_dinh')
    ws_qd.views.sheetView[0].showGridLines = True

    ws_qd['A1'] = "BẢNG QUYẾT ĐỊNH SỐ QUẦY THEO KHUNG GIỜ (SIPP LÝ THUYẾT - ERLANG C VÀ ERLANG-A)"
    ws_qd['A1'].font = font_title
    ws_qd['A2'] = (f"Tham số tập train (10 ngày): E[S] = {es_train:.4f} phút, cs² = {cs2_train:.4f}, "
                   f"θ = {theta_train:.4f}/phút. SLA: P(phục vụ và chờ <= 5 phút) >= 90%")
    ws_qd['A2'].font = font_bold

    headers_qd = [
        "Khung giờ", "λ train (khách/phút)", "E[S] (phút)", "Tải a (Erlang)",
        "Số quầy c (Erlang C)", "SL Erlang C", "Số quầy c (Erlang-A)", "SL Erlang-A",
        "p_ab Erlang-A", "Hiệu suất ρ (Erlang-A)"
    ]
    for idx, h_text in enumerate(headers_qd, start=1):
        cell = ws_qd.cell(row=4, column=idx, value=h_text)
        cell.font = font_header
        cell.fill = fill_header_dark
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for idx, lam in enumerate(train_lambdas, start=5):
        h_label = f"{idx+3}h - {idx+4}h"
        c_c = bt.sipp_schedule([lam], mu_train, theta_train, use_erlang_a=False)[0]
        c_a = bt.sipp_schedule([lam], mu_train, theta_train, use_erlang_a=True)[0]
        res_c = bt.mmc(lam, mu_train, c_c)
        res_a = bt.erlang_a(lam, mu_train, theta_train, c_a)

        ws_qd[f'A{idx}'] = h_label
        ws_qd[f'B{idx}'] = lam
        ws_qd[f'C{idx}'] = es_train
        ws_qd[f'D{idx}'] = f'=B{idx}*C{idx}'
        ws_qd[f'E{idx}'] = c_c
        ws_qd[f'F{idx}'] = round(res_c["P_sla"], 4)
        ws_qd[f'G{idx}'] = c_a
        ws_qd[f'H{idx}'] = round(res_a["SL"], 4)
        ws_qd[f'I{idx}'] = round(res_a["p_ab"], 4)
        ws_qd[f'J{idx}'] = round(res_a["util"], 4)

        ws_qd[f'A{idx}'].font = font_bold
        ws_qd[f'B{idx}'].number_format = '0.0000'
        ws_qd[f'C{idx}'].number_format = '0.0000'
        ws_qd[f'D{idx}'].number_format = '0.00'
        for col_l in ['F', 'H', 'I', 'J']:
            ws_qd[f'{col_l}{idx}'].number_format = '0.0%'
        for col_l in ['E', 'G']:
            ws_qd[f'{col_l}{idx}'].font = Font(name='Segoe UI', size=11, bold=True, color='006100')
            ws_qd[f'{col_l}{idx}'].fill = fill_pass
        for col_l in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J']:
            ws_qd[f'{col_l}{idx}'].border = border_thin
            ws_qd[f'{col_l}{idx}'].alignment = Alignment(horizontal='center')

    ws_qd['A14'] = "TỔNG QUẦY-GIỜ"
    ws_qd['A14'].font = font_bold
    ws_qd['A14'].alignment = Alignment(horizontal='center')
    for col_l in ['E', 'G']:
        ws_qd[f'{col_l}14'] = f"=SUM({col_l}5:{col_l}13)"
        ws_qd[f'{col_l}14'].font = Font(name='Segoe UI', size=12, bold=True, color='1F4E78')
        ws_qd[f'{col_l}14'].alignment = Alignment(horizontal='center')

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(excel_path)
    print(f"Xuat file Excel Erlang thanh cong ra: {excel_path}")

if __name__ == '__main__':
    build_erlang_excel()
